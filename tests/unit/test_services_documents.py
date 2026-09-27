# ruff: noqa: E501

from __future__ import annotations

import io
import zipfile

import fitz
import pytest
from docx import Document
from openpyxl import Workbook

from app.services.documents import DocumentExtractionStatus, DocumentExtractor
from app.services.documents.ocr import DisabledOCRProvider
from app.services.documents.security import ArchiveLimits, UnsafeArchiveError, read_safe_zip


def test_extracts_plain_html_and_csv_with_offsets() -> None:
    extractor = DocumentExtractor()
    plain = extractor.extract("Razões recursais\nEmpresa ABC".encode(), filename="ata.txt")
    assert plain.status is DocumentExtractionStatus.EXTRACTED
    assert plain.chunks[0].end_offset == len(plain.text)

    html = extractor.extract(
        b"<html><script>secret()</script><body><h1>Decisao</h1><p>Recurso deferido</p></body></html>",
        filename="decisao.html",
    )
    assert "Recurso deferido" in html.text
    assert "secret" not in html.text

    csv_result = extractor.extract(b"empresa;status\nABC;inabilitada", filename="resultado.csv")
    assert csv_result.chunks[1].locator == "row:2"
    assert "ABC | inabilitada" in csv_result.text


def test_extracts_docx_and_xlsx_locations() -> None:
    document = Document()
    document.add_paragraph("Empresa ABC foi inabilitada")
    docx_bytes = io.BytesIO()
    document.save(docx_bytes)
    docx = DocumentExtractor().extract(docx_bytes.getvalue(), filename="ata.docx")
    assert docx.status is DocumentExtractionStatus.EXTRACTED
    assert docx.chunks[0].locator == "paragraph:1"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Resultado"
    sheet.append(["Empresa", "Situação"])
    sheet.append(["ABC", "Desclassificada"])
    xlsx_bytes = io.BytesIO()
    workbook.save(xlsx_bytes)
    xlsx = DocumentExtractor().extract(xlsx_bytes.getvalue(), filename="resultado.xlsx")
    assert xlsx.chunks[0].metadata["sheet"] == "Resultado"
    assert xlsx.chunks[0].metadata["cell_range"] == "A1:B2"


def test_pdf_pages_and_scanned_pdf_status() -> None:
    textual = fitz.open()
    page = textual.new_page()
    page.insert_text((72, 72), "Ata de julgamento com texto suficiente para extracao")
    result = DocumentExtractor().extract(textual.tobytes(), filename="ata.pdf")
    textual.close()
    assert result.status is DocumentExtractionStatus.EXTRACTED
    assert result.chunks[0].page_number == 1

    image_pdf = fitz.open()
    image_page = image_pdf.new_page()
    pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 10, 10), False)
    pixmap.clear_with(255)
    image_page.insert_image(fitz.Rect(0, 0, 10, 10), pixmap=pixmap)
    scanned = DocumentExtractor().extract(image_pdf.tobytes(), filename="scan.pdf")
    image_pdf.close()
    assert scanned.status is DocumentExtractionStatus.OCR_REQUIRED


def test_zip_is_extracted_in_memory_and_rejects_traversal() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("atas/ata.txt", "licitante inabilitada")
        archive.writestr("resultado.csv", "empresa,status\nABC,vencedora")
    result = DocumentExtractor().extract(buffer.getvalue(), filename="documentos.zip")
    assert result.status is DocumentExtractionStatus.EXTRACTED
    assert {member.filename for member in result.archive_members} == {
        "atas/ata.txt",
        "resultado.csv",
    }

    unsafe = io.BytesIO()
    with zipfile.ZipFile(unsafe, "w") as archive:
        archive.writestr("../escape.txt", "x")
    with pytest.raises(UnsafeArchiveError):
        read_safe_zip(unsafe.getvalue())


def test_zip_limits_and_disabled_ocr() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("one.txt", "a")
        archive.writestr("two.txt", "b")
    with pytest.raises(UnsafeArchiveError, match="too many"):
        read_safe_zip(buffer.getvalue(), ArchiveLimits(max_members=1))

    assert (
        __import__("asyncio").run(DisabledOCRProvider().extract_text(b"x", mime_type="image/png"))
        == []
    )
