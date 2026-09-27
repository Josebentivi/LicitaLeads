"""Deterministic, bounded extraction for supported public-procurement files."""

# ruff: noqa: E501

from __future__ import annotations

import csv
import hashlib
import html
import io
import mimetypes
import re
from collections.abc import Iterable
from pathlib import PurePath

from .security import ArchiveLimits, UnsafeArchiveError, read_safe_zip
from .types import (
    DocumentChunk,
    DocumentExtraction,
    DocumentExtractionStatus,
    ExtractedArchiveMember,
)

_WHITESPACE = re.compile(r"[ \t\f\v]+")
_BLANK_LINES = re.compile(r"\n{3,}")
_SCRIPT_STYLE = re.compile(r"<(script|style|noscript)\b[^>]*>.*?</\1>", re.I | re.S)
_HTML_TAG = re.compile(r"<[^>]+>")

_EXTENSION_MIME = {
    ".pdf": "application/pdf",
    ".html": "text/html",
    ".htm": "text/html",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
    ".txt": "text/plain",
    ".zip": "application/zip",
}


def _clean_text(value: str) -> str:
    lines = [_WHITESPACE.sub(" ", line).strip() for line in value.replace("\r", "\n").split("\n")]
    return _BLANK_LINES.sub("\n\n", "\n".join(line for line in lines if line)).strip()


def _decode_text(content: bytes) -> str:
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        return content.decode("utf-16")
    if content.startswith(b"\xef\xbb\xbf"):
        return content.decode("utf-8-sig")
    for encoding in ("utf-8", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _offset_chunks(
    items: Iterable[tuple[str, int | None, str | None, dict]],
) -> tuple[str, tuple[DocumentChunk, ...]]:
    text_parts: list[str] = []
    chunks: list[DocumentChunk] = []
    cursor = 0
    for raw_text, page, locator, metadata in items:
        chunk_text = _clean_text(raw_text)
        if not chunk_text:
            continue
        if text_parts:
            cursor += 2
        start = cursor
        text_parts.append(chunk_text)
        cursor += len(chunk_text)
        chunks.append(
            DocumentChunk(
                text=chunk_text,
                page_number=page,
                locator=locator,
                start_offset=start,
                end_offset=cursor,
                metadata=metadata,
            )
        )
    return "\n\n".join(text_parts), tuple(chunks)


class DocumentExtractor:
    """Extract supported formats without writing untrusted paths to disk."""

    def __init__(
        self,
        *,
        archive_limits: ArchiveLimits | None = None,
        min_pdf_text_characters: int = 20,
        max_nested_archives: int = 1,
    ) -> None:
        self.archive_limits = archive_limits or ArchiveLimits()
        self.min_pdf_text_characters = min_pdf_text_characters
        self.max_nested_archives = max_nested_archives

    def extract(
        self,
        content: bytes,
        *,
        filename: str,
        mime_type: str | None = None,
        _archive_depth: int = 0,
    ) -> DocumentExtraction:
        digest = hashlib.sha256(content).hexdigest()
        detected = self.detect_mime(content, filename)
        supplied = mime_type.split(";", 1)[0].strip().lower() if mime_type else None
        warnings: list[str] = []
        if supplied and supplied not in {detected, "application/octet-stream"}:
            warnings.append(f"declared MIME {supplied!r} differs from detected MIME {detected!r}")
        try:
            if detected == "application/pdf":
                return self._pdf(content, detected, digest, warnings)
            if detected == "text/html":
                return self._html(content, detected, digest, warnings)
            if detected.endswith("wordprocessingml.document"):
                return self._docx(content, detected, digest, warnings)
            if detected.endswith("spreadsheetml.sheet"):
                return self._xlsx(content, detected, digest, warnings)
            if detected == "text/csv":
                return self._csv(content, detected, digest, warnings)
            if detected == "text/plain":
                return self._plain(content, detected, digest, warnings)
            if detected == "application/zip":
                return self._zip(content, detected, digest, warnings, _archive_depth)
            return DocumentExtraction(
                status=DocumentExtractionStatus.UNSUPPORTED,
                mime_type=detected,
                sha256=digest,
                warnings=tuple(warnings),
            )
        except (UnsafeArchiveError, ValueError):
            raise
        except Exception as exc:  # malformed third-party file, kept as an auditable failure
            return DocumentExtraction(
                status=DocumentExtractionStatus.FAILED,
                mime_type=detected,
                sha256=digest,
                warnings=tuple([*warnings, f"{type(exc).__name__}: {exc}"]),
            )

    @staticmethod
    def detect_mime(content: bytes, filename: str) -> str:
        extension = PurePath(filename).suffix.lower()
        if content.startswith(b"%PDF-"):
            return "application/pdf"
        # OOXML documents are ZIP containers, so their extension is material.
        if content.startswith(b"PK\x03\x04") and extension in {".docx", ".xlsx"}:
            return _EXTENSION_MIME[extension]
        if content.startswith(b"PK\x03\x04"):
            return "application/zip"
        sample = content[:1024].lstrip().lower()
        if extension in {".html", ".htm"} or sample.startswith((b"<!doctype html", b"<html")):
            return "text/html"
        if extension in _EXTENSION_MIME:
            return _EXTENSION_MIME[extension]
        guessed, _ = mimetypes.guess_type(filename)
        return guessed or "application/octet-stream"

    def _pdf(
        self, content: bytes, mime: str, digest: str, warnings: list[str]
    ) -> DocumentExtraction:
        try:
            import fitz  # type: ignore[import-untyped]
        except ImportError as exc:  # pragma: no cover - dependency declared by the project
            raise RuntimeError("PyMuPDF is required to extract PDF documents") from exc
        items: list[tuple[str, int | None, str | None, dict]] = []
        image_count = 0
        with fitz.open(stream=content, filetype="pdf") as document:
            for index, page in enumerate(document):
                page_text = page.get_text("text")
                images = len(page.get_images(full=True))
                image_count += images
                items.append((page_text, index + 1, f"page:{index + 1}", {"image_count": images}))
        text, chunks = _offset_chunks(items)
        if len(text.strip()) < self.min_pdf_text_characters and image_count:
            return DocumentExtraction(
                status=DocumentExtractionStatus.OCR_REQUIRED,
                mime_type=mime,
                sha256=digest,
                text=text,
                chunks=chunks,
                warnings=tuple(warnings),
            )
        status = DocumentExtractionStatus.EXTRACTED if text else DocumentExtractionStatus.EMPTY
        return DocumentExtraction(status, mime, digest, text, chunks, warnings=tuple(warnings))

    def _html(
        self, content: bytes, mime: str, digest: str, warnings: list[str]
    ) -> DocumentExtraction:
        source = _decode_text(content)
        try:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(source, "html.parser")
            for unwanted in soup(["script", "style", "noscript", "template"]):
                unwanted.decompose()
            extracted = soup.get_text("\n")
        except ImportError:  # pragma: no cover - defensive fallback
            extracted = _HTML_TAG.sub("\n", _SCRIPT_STYLE.sub("", source))
            extracted = html.unescape(extracted)
        text, chunks = _offset_chunks([(extracted, None, "html:body", {})])
        status = DocumentExtractionStatus.EXTRACTED if text else DocumentExtractionStatus.EMPTY
        return DocumentExtraction(status, mime, digest, text, chunks, warnings=tuple(warnings))

    def _docx(
        self, content: bytes, mime: str, digest: str, warnings: list[str]
    ) -> DocumentExtraction:
        from docx import Document

        document = Document(io.BytesIO(content))
        items: list[tuple[str, int | None, str | None, dict]] = []
        for index, paragraph in enumerate(document.paragraphs, start=1):
            items.append((paragraph.text, None, f"paragraph:{index}", {}))
        for table_index, table in enumerate(document.tables, start=1):
            rows = [" | ".join(cell.text for cell in row.cells) for row in table.rows]
            items.append(("\n".join(rows), None, f"table:{table_index}", {}))
        text, chunks = _offset_chunks(items)
        status = DocumentExtractionStatus.EXTRACTED if text else DocumentExtractionStatus.EMPTY
        return DocumentExtraction(status, mime, digest, text, chunks, warnings=tuple(warnings))

    def _xlsx(
        self, content: bytes, mime: str, digest: str, warnings: list[str]
    ) -> DocumentExtraction:
        from openpyxl import load_workbook  # type: ignore[import-untyped]
        from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        items: list[tuple[str, int | None, str | None, dict]] = []
        try:
            for sheet in workbook.worksheets:
                rows: list[str] = []
                max_column = 0
                last_row = 0
                for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    values = ["" if value is None else str(value) for value in row]
                    while values and values[-1] == "":
                        values.pop()
                    if values:
                        rows.append(" | ".join(values))
                        max_column = max(max_column, len(values))
                        last_row = row_index
                cell_range = f"A1:{get_column_letter(max_column)}{last_row}" if rows else None
                items.append(
                    (
                        "\n".join(rows),
                        None,
                        f"sheet:{sheet.title}",
                        {"sheet": sheet.title, "cell_range": cell_range},
                    )
                )
        finally:
            workbook.close()
        text, chunks = _offset_chunks(items)
        status = DocumentExtractionStatus.EXTRACTED if text else DocumentExtractionStatus.EMPTY
        return DocumentExtraction(status, mime, digest, text, chunks, warnings=tuple(warnings))

    def _csv(
        self, content: bytes, mime: str, digest: str, warnings: list[str]
    ) -> DocumentExtraction:
        source = _decode_text(content)
        try:
            dialect = csv.Sniffer().sniff(source[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [
            " | ".join(cell.strip() for cell in row)
            for row in csv.reader(io.StringIO(source), dialect)
        ]
        text, chunks = _offset_chunks(
            [(row, None, f"row:{index}", {"row": index}) for index, row in enumerate(rows, start=1)]
        )
        status = DocumentExtractionStatus.EXTRACTED if text else DocumentExtractionStatus.EMPTY
        return DocumentExtraction(status, mime, digest, text, chunks, warnings=tuple(warnings))

    def _plain(
        self, content: bytes, mime: str, digest: str, warnings: list[str]
    ) -> DocumentExtraction:
        text, chunks = _offset_chunks([(_decode_text(content), None, "text", {})])
        status = DocumentExtractionStatus.EXTRACTED if text else DocumentExtractionStatus.EMPTY
        return DocumentExtraction(status, mime, digest, text, chunks, warnings=tuple(warnings))

    def _zip(
        self,
        content: bytes,
        mime: str,
        digest: str,
        warnings: list[str],
        depth: int,
    ) -> DocumentExtraction:
        if depth >= self.max_nested_archives:
            raise UnsafeArchiveError("nested archive depth exceeded")
        members: list[ExtractedArchiveMember] = []
        for filename, member_content in read_safe_zip(content, self.archive_limits):
            extraction = self.extract(
                member_content,
                filename=filename,
                _archive_depth=depth + 1,
            )
            members.append(ExtractedArchiveMember(filename, extraction))
        combined_items: list[tuple[str, int | None, str | None, dict]] = []
        for archive_member in members:
            if archive_member.extraction.text:
                combined_items.append(
                    (
                        archive_member.extraction.text,
                        None,
                        f"archive:{archive_member.filename}",
                        {
                            "archive_member": archive_member.filename,
                            "sha256": archive_member.extraction.sha256,
                        },
                    )
                )
        text, chunks = _offset_chunks(combined_items)
        status = DocumentExtractionStatus.EXTRACTED if text else DocumentExtractionStatus.EMPTY
        return DocumentExtraction(
            status=status,
            mime_type=mime,
            sha256=digest,
            text=text,
            chunks=chunks,
            archive_members=tuple(members),
            warnings=tuple(warnings),
        )
